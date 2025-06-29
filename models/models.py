# -*- coding: utf-8 -*-

from odoo import models, fields, api
import requests
import base64
from io import BytesIO
import openpyxl

WHATSAPP_API_URL = f"https://graph.facebook.com/v19.0/337320726121483/messages"
ACCESS_TOKEN = "Bearer EAAO5XN3WP2YBO4CSvZAFfh0k7ZCHsOi9ck5v2hDMZBYK0wIV0ZABizqmWoZAd8yYwrw2VgeMEE789aBDwiCBieiZCYta0YJzTFbkyG8a3glpvZBkRzNxXiIAw5sZASDrinWU4aLtt5NGnFIdDCoELdZBMPQeEpKgXgXCsG6Tj0SiV717k0Y4zGJJyiSt3m5zfeEu1xMZBYcPtmQQZBteTik4WQWA4iDjhuZA0dCyFlGxMe4q"



class ConsultingDocument(models.Model):
    _name = 'consulting.document'
    _description = 'Consulting Document'

    name = fields.Char(string="Document Name", required=True)
    document_file = fields.Binary(string="File", required=True)
    location = fields.Char(string="Location")
    uploaded_by = fields.Many2one('hr.employee', string="Uploaded By")
    upload_date = fields.Datetime(string="Upload Date", default=fields.Datetime.now)
    whatsapp_message_log_ids = fields.One2many(
        'consulting.whatsapp.message.log',
        'consulting_document_id',
        string="WhatsApp Message Logs"
    )


    def create_whatsapp_message_log(self, name, mobile, location, qualification):
        self.ensure_one()
        headers = {
            "Authorization": f"{ACCESS_TOKEN}",
            "Content-Type": "application/json",
        }
        mobile = f"91{mobile}"
        data = {
            "messaging_product": "whatsapp",
            "to": mobile,
            "type": "template",
            "template": {
                "name": "test_user_accept",
                "language": {
                    "code": "en"
                },
                "components": [
                    {
                        "type": "body",
                        "parameters": [
                            {
                                "type": "text",
                                "text": "demo user" 
                            }
                        ]
                    }
                ]
            }
        }
        try:
            response = requests.post(WHATSAPP_API_URL, json=data, headers=headers)
            status_code = response.status_code
            response_text = response.text
        except Exception as e:
            status_code = 0
            response_text = str(e)

        return self.env['consulting.whatsapp.message.log'].create({
            'full_name': name,
            'phone_number': mobile,
            'location': location,
            'qualification': qualification,
            'status_code': status_code,
            'response_text': response_text,
            'consulting_document_id': self.id,
        })

    def action_import_and_send_whatsapp(self):
        self.ensure_one()
        if not self.document_file:
            raise ValueError("No file attached to this document.")

        # Decode and load the XLSX file
        file_content = base64.b64decode(self.document_file)
        workbook = openpyxl.load_workbook(BytesIO(file_content))
        sheet = workbook.active

        # Get headers from the first row
        headers = [cell.value for cell in next(sheet.iter_rows(min_row=1, max_row=1))]
        # Map header names to column indices
        header_map = {header.lower(): idx for idx, header in enumerate(headers)}

        for row in sheet.iter_rows(min_row=2, values_only=True):
            name = row[header_map.get('full name')]
            mobile = row[header_map.get('mobile no')]
            location = row[header_map.get('location')]
            qualification = row[header_map.get('qualification')]
            if name and mobile:
                self.create_whatsapp_message_log(name, mobile, location, qualification)



class ConsultingWhatsAppMessageLog(models.Model):
    _name = 'consulting.whatsapp.message.log'
    _description = 'WhatsApp Message Log'

    phone_number = fields.Char(string="Phone Number", required=True)
    full_name = fields.Char(string="Full Name")
    location = fields.Char(string="Location")
    qualification = fields.Char(string="Qualification")
    status_code = fields.Integer(string="Status Code")
    response_text = fields.Text(string="Response Text")
    sent_at = fields.Datetime(string="Sent At")
    consulting_document_id = fields.Many2one(
        'consulting.document',
        string="Consulting Document",
        ondelete='cascade'
    )
    reply_text = fields.Char(string="Reply Text")
    reply_date = fields.Datetime(string="Reply Date")